import openpyxl
import re
import pandas as pd
import os

file_name = '2023-11-30_insert_member.xlsx'
output_path = f"{file_name}"

wb = openpyxl.load_workbook(file_name)
sheet = wb.active

info_warning = False
phone_pattern = r'\d{3}-\d{3,4}-\d{4}'
email_pattern = r"[a-zA-Z0-9._+-]+@[a-zA-Z0-9]+\.[a-zA-Z]{2,4}"
for row in sheet.iter_rows():
    for cell in row:
        if re.findall(phone_pattern,str(cell.value)):
            info_warning = True
        elif re.findall(email_pattern,str(cell.value)):
            info_warning = True
if info_warning:
    print("위험 정보 있음")
else:
    print("위험 정보 없음")

safe_warning = False
#for row in sheet.iter_rows():
#    print(row[6:11])
    #for cell in row[7]:
        #print()
        #if re.findall('False',str(cell.value)):
        #    safe_warning = True
        #elif safe_warning:
        #    for cell in row[6:11]:
        #        if re.findall('False',str(cell.value)):
        
df = pd.read_excel(file_name)
df.to_csv(f'{file_name}.txt',sep='|',index=False)

with open(f'{file_name}.txt','r',encoding='utf-8') as f:
    lines = f.readlines()
    head_info=list(map(str,lines[0].split('|')))
    print(head_info)
    print()
    for line in lines[1:]:
        print(line)

